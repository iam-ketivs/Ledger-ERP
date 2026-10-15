"""Report building and exporting.

One place that knows how to assemble every report in Ledger, so CSV, Excel and the
printable view all show exactly the same numbers. Adding a report here makes it
exportable in all three formats at once.

Each report returns:
    title    - heading
    subtitle - what period / filter it covers
    columns  - [(label, kind)] where kind is text | money | number | date
    rows     - list of lists, matching columns
    totals   - optional list matching columns, for the bottom line
"""
import csv
import io
from datetime import date, datetime


def _d(s):
    return datetime.strptime(s, "%Y-%m-%d").date()


def _range(args, default_from=None, default_to=None):
    today = date.today()
    f = args.get("from") or (default_from.isoformat() if default_from else "")
    t = args.get("to") or (default_to.isoformat() if default_to else today.isoformat())
    try:
        start = _d(f) if f else None
    except ValueError:
        start = None
    try:
        end = _d(t) if t else today
    except ValueError:
        end = today
    return start, end


def _period_label(start, end):
    if not start:
        return f"All time up to {end.strftime('%d %b %Y')}"
    if start == end:
        return start.strftime("%d %b %Y")
    return f"{start.strftime('%d %b %Y')} to {end.strftime('%d %b %Y')}"


# --------------------------------------------------------------------------- reports

def report_cashbook(app, conn, args):
    start, end = _range(args)
    moves = app.cash_movements(conn, start, end)
    opening = 0.0
    if start:
        from datetime import timedelta
        prev = app.cash_movements(conn, None, start - timedelta(days=1))
        opening = sum((m["amount"] if m["direction"] == "in" else -m["amount"]) for m in prev)
    running = opening
    rows = [["", "Opening balance", "", None, None, opening]]
    tin = tout = 0.0
    for m in moves:
        if m["direction"] == "in":
            running += m["amount"]; tin += m["amount"]
            rows.append([m["date"], m["label"], m["detail"], m["amount"], None, running])
        else:
            running -= m["amount"]; tout += m["amount"]
            rows.append([m["date"], m["label"], m["detail"], None, m["amount"], running])
    return {
        "title": "Cash Book",
        "subtitle": _period_label(start, end),
        "columns": [("Date", "date"), ("Description", "text"), ("Detail", "text"),
                    ("Money in", "money"), ("Money out", "money"), ("Balance", "money")],
        "rows": rows,
        "totals": ["", "Closing balance", "", tin, tout, running],
    }


def report_sales(app, conn, args):
    start, end = _range(args)
    lo = start.isoformat() if start else "0000-01-01"
    q = ("SELECT s.*, c.name cname FROM sales s LEFT JOIN customers c ON c.id=s.customer_id "
         "WHERE s.sale_date BETWEEN ? AND ? ORDER BY s.sale_date, s.id")
    rows = []
    tt = tc = tp = tpaid = tdisc = 0.0
    for r in conn.execute(q, (lo, end.isoformat())):
        rows.append([r["sale_date"], f"#{r['id']}", r["cname"] or "Walk-in",
                     r["discount_amount"] or 0, r["total_amount"], r["total_cost"],
                     r["total_profit"], r["paid_amount"], r["payment_status"], r["notes"] or ""])
        tdisc += r["discount_amount"] or 0; tt += r["total_amount"]; tc += r["total_cost"]
        tp += r["total_profit"]; tpaid += r["paid_amount"]
    return {
        "title": "Sales Report",
        "subtitle": _period_label(start, end),
        "columns": [("Date", "date"), ("Sale", "text"), ("Customer", "text"),
                    ("Discount", "money"), ("Total", "money"), ("Cost", "money"),
                    ("Profit", "money"), ("Paid", "money"), ("Status", "text"), ("Notes", "text")],
        "rows": rows,
        "totals": ["", "", f"{len(rows)} sales", tdisc, tt, tc, tp, tpaid, "", ""],
    }


def report_sale_items(app, conn, args):
    start, end = _range(args)
    lo = start.isoformat() if start else "0000-01-01"
    rows = []
    tq = tt = tp = 0.0
    for r in conn.execute(
        "SELECT s.sale_date, s.id sid, c.name cname, i.name iname, i.unit, si.quantity, "
        "si.unit_price, si.unit_cost, si.line_total, si.line_profit "
        "FROM sale_items si JOIN sales s ON s.id=si.sale_id JOIN items i ON i.id=si.item_id "
        "LEFT JOIN customers c ON c.id=s.customer_id "
        "WHERE s.sale_date BETWEEN ? AND ? ORDER BY s.sale_date, s.id", (lo, end.isoformat())):
        rows.append([r["sale_date"], f"#{r['sid']}", r["cname"] or "Walk-in", r["iname"],
                     r["quantity"], r["unit"], r["unit_price"], r["unit_cost"],
                     r["line_total"], r["line_profit"]])
        tq += r["quantity"]; tt += r["line_total"]; tp += r["line_profit"]
    return {
        "title": "Sales - Item Detail",
        "subtitle": _period_label(start, end),
        "columns": [("Date", "date"), ("Sale", "text"), ("Customer", "text"), ("Item", "text"),
                    ("Qty", "number"), ("Unit", "text"), ("Price", "money"), ("Cost", "money"),
                    ("Line total", "money"), ("Line profit", "money")],
        "rows": rows,
        "totals": ["", "", "", f"{len(rows)} lines", tq, "", "", "", tt, tp],
    }


def report_purchases(app, conn, args):
    start, end = _range(args)
    lo = start.isoformat() if start else "0000-01-01"
    rows = []
    tt = tpaid = 0.0
    for r in conn.execute(
        "SELECT p.*, i.name iname, i.unit, s.name sname FROM purchases p "
        "JOIN items i ON i.id=p.item_id LEFT JOIN suppliers s ON s.id=p.supplier_id "
        "WHERE p.purchase_date BETWEEN ? AND ? ORDER BY p.purchase_date, p.id",
            (lo, end.isoformat())):
        owed = r["total_cost"] - (r["paid_amount"] or 0)
        rows.append([r["purchase_date"], r["iname"], r["sname"] or r["supplier_name"] or "",
                     r["quantity"], r["unit"], r["unit_cost"], r["total_cost"],
                     r["paid_amount"] or 0, owed, r["payment_status"], r["notes"] or ""])
        tt += r["total_cost"]; tpaid += r["paid_amount"] or 0
    return {
        "title": "Purchase Report",
        "subtitle": _period_label(start, end),
        "columns": [("Date", "date"), ("Item", "text"), ("Supplier", "text"), ("Qty", "number"),
                    ("Unit", "text"), ("Unit cost", "money"), ("Total", "money"),
                    ("Paid", "money"), ("Still owed", "money"), ("Status", "text"), ("Notes", "text")],
        "rows": rows,
        "totals": ["", "", f"{len(rows)} purchases", "", "", "", tt, tpaid, tt - tpaid, "", ""],
    }


def report_orders(app, conn, args):
    start, end = _range(args)
    lo = start.isoformat() if start else "0000-01-01"
    rows = []
    tt = tadv = tprof = 0.0
    for r in conn.execute(
        "SELECT o.*, c.name cname FROM orders o JOIN customers c ON c.id=o.customer_id "
        "WHERE o.order_date BETWEEN ? AND ? ORDER BY o.order_date, o.id", (lo, end.isoformat())):
        rows.append([r["order_date"], f"#{r['id']}", r["cname"], r["agreed_total"],
                     r["advance_amount"], r["balance_received"] or 0,
                     r["refunded_amount"] or 0, r["total_cost"] or 0, r["profit"] or 0,
                     r["status"], r["completed_date"] or r["cancelled_date"] or ""])
        tt += r["agreed_total"]; tadv += r["advance_amount"]; tprof += r["profit"] or 0
    return {
        "title": "Orders Report",
        "subtitle": _period_label(start, end),
        "columns": [("Booked", "date"), ("Order", "text"), ("Customer", "text"),
                    ("Agreed total", "money"), ("Advance", "money"), ("Balance recd", "money"),
                    ("Refunded", "money"), ("Cost", "money"), ("Profit", "money"),
                    ("Status", "text"), ("Closed on", "date")],
        "rows": rows,
        "totals": ["", "", f"{len(rows)} orders", tt, tadv, "", "", "", tprof, "", ""],
    }


def report_customers(app, conn, args):
    rows = []
    tc = 0.0
    for r in conn.execute(
        "SELECT c.*, (SELECT COUNT(*) FROM sales s WHERE s.customer_id=c.id) nsales, "
        "(SELECT COALESCE(SUM(total_amount),0) FROM sales s WHERE s.customer_id=c.id) spend "
        "FROM customers c WHERE c.is_active=1 ORDER BY c.name"):
        rows.append([r["name"], r["phone"] or "", r["address"] or "",
                     "Regular" if r["customer_type"] == "regular" else "One-time",
                     r["nsales"], r["spend"], r["credit_balance"]])
        tc += r["credit_balance"]
    return {
        "title": "Customer Report",
        "subtitle": f"All active customers as at {date.today().strftime('%d %b %Y')}",
        "columns": [("Name", "text"), ("Phone", "text"), ("Address", "text"), ("Type", "text"),
                    ("Sales", "number"), ("Total bought", "money"), ("Owes you", "money")],
        "rows": rows,
        "totals": [f"{len(rows)} customers", "", "", "", "", "", tc],
    }


def report_suppliers(app, conn, args):
    rows = []
    tp = 0.0
    for r in conn.execute(
        "SELECT s.*, (SELECT COUNT(*) FROM purchases p WHERE p.supplier_id=s.id) n, "
        "(SELECT COALESCE(SUM(total_cost),0) FROM purchases p WHERE p.supplier_id=s.id) spend "
        "FROM suppliers s WHERE s.is_active=1 ORDER BY s.name"):
        rows.append([r["name"], r["phone"] or "", r["address"] or "", r["n"], r["spend"],
                     r["payable_balance"]])
        tp += r["payable_balance"]
    return {
        "title": "Supplier Report",
        "subtitle": f"All active suppliers as at {date.today().strftime('%d %b %Y')}",
        "columns": [("Name", "text"), ("Phone", "text"), ("Address", "text"),
                    ("Purchases", "number"), ("Total bought", "money"), ("You owe", "money")],
        "rows": rows,
        "totals": [f"{len(rows)} suppliers", "", "", "", "", tp],
    }


def report_items(app, conn, args):
    rows = []
    tv = 0.0
    for r in conn.execute("SELECT * FROM items WHERE is_active=1 ORDER BY item_type, name"):
        value = (r["stock_qty"] or 0) * (r["avg_cost"] or 0)
        rows.append([r["name"], r["item_type"].replace("_", " "), r["category"] or "",
                     r["unit"], r["stock_qty"], r["avg_cost"], value, r["selling_price"],
                     r["reorder_level"]])
        if r["item_type"] != "service":
            tv += value
    return {
        "title": "Stock Report",
        "subtitle": f"All active items as at {date.today().strftime('%d %b %Y')}",
        "columns": [("Item", "text"), ("Type", "text"), ("Category", "text"), ("Unit", "text"),
                    ("Stock", "number"), ("Avg cost", "money"), ("Stock value", "money"),
                    ("Selling price", "money"), ("Reorder at", "number")],
        "rows": rows,
        "totals": [f"{len(rows)} items", "", "", "", "", "", tv, "", ""],
    }


def report_expenses(app, conn, args):
    start, end = _range(args)
    lo = start.isoformat() if start else "0000-01-01"
    rows = []
    t = 0.0
    for r in conn.execute(
        "SELECT * FROM expenses WHERE expense_date BETWEEN ? AND ? ORDER BY expense_date, id",
            (lo, end.isoformat())):
        rows.append([r["expense_date"], r["category"] or "", r["amount"], r["notes"] or ""])
        t += r["amount"]
    return {
        "title": "Expense Report",
        "subtitle": _period_label(start, end),
        "columns": [("Date", "date"), ("Category", "text"), ("Amount", "money"), ("Notes", "text")],
        "rows": rows,
        "totals": ["", f"{len(rows)} entries", t, ""],
    }


def report_production(app, conn, args):
    start, end = _range(args)
    lo = start.isoformat() if start else "0000-01-01"
    rows = []
    tc = 0.0
    for r in conn.execute(
        "SELECT pr.*, i.name iname, i.unit FROM production_runs pr "
        "JOIN items i ON i.id=pr.output_item_id "
        "WHERE pr.production_date BETWEEN ? AND ? ORDER BY pr.production_date, pr.id",
            (lo, end.isoformat())):
        total = (r["total_material_cost"] or 0) + (r["labor_cost"] or 0) + (r["other_cost"] or 0)
        rows.append([r["production_date"], r["iname"], r["output_qty"], r["unit"],
                     r["total_material_cost"] or 0, r["labor_cost"] or 0, r["other_cost"] or 0,
                     total, r["cost_per_unit"], r["notes"] or ""])
        tc += total
    return {
        "title": "Production Report",
        "subtitle": _period_label(start, end),
        "columns": [("Date", "date"), ("Output item", "text"), ("Qty made", "number"),
                    ("Unit", "text"), ("Material", "money"), ("Labour", "money"),
                    ("Other", "money"), ("Total cost", "money"), ("Cost/unit", "money"),
                    ("Notes", "text")],
        "rows": rows,
        "totals": ["", f"{len(rows)} runs", "", "", "", "", "", tc, "", ""],
    }


def report_employees(app, conn, args):
    rows = []
    ts = 0.0
    for r in conn.execute("SELECT * FROM employees WHERE is_active=1 ORDER BY name"):
        adv = conn.execute(
            "SELECT COALESCE(SUM(amount),0) t FROM employee_advances WHERE employee_id=?",
            (r["id"],)).fetchone()["t"]
        rows.append([r["name"], r["phone"] or "", r["monthly_salary"], adv,
                     (r["created_at"] or "")[:10]])
        ts += r["monthly_salary"]
    return {
        "title": "Employee Report",
        "subtitle": f"All active staff as at {date.today().strftime('%d %b %Y')}",
        "columns": [("Name", "text"), ("Phone", "text"), ("Monthly salary", "money"),
                    ("Advances (all time)", "money"), ("Added on", "date")],
        "rows": rows,
        "totals": [f"{len(rows)} staff", "", ts, "", ""],
    }


def report_summary(app, conn, args):
    """The overall picture - trading result plus where the money sits."""
    start, end = _range(args)
    sales, all_orders, comp = app.load_revenue_sources(conn)
    st = app.period_stats(sales, comp, start, end)
    cash = app.period_cash(conn, sales, all_orders, comp, start, end)
    pos = app.money_position(conn, sales, all_orders, comp)
    lo = start.isoformat() if start else "0000-01-01"
    exp = conn.execute(
        "SELECT COALESCE(SUM(amount),0) t FROM expenses WHERE expense_date BETWEEN ? AND ?",
        (lo, end.isoformat())).fetchone()["t"]

    rows = [
        ["TRADING", "", ""],
        ["Revenue", st["revenue"], "sales + custom orders completed"],
        ["Cost of goods sold", -st["cost"], "material cost of what was sold"],
        ["Gross profit", st["profit"], f"margin {st['revenue'] and st['profit']/st['revenue']*100 or 0:.1f}%"],
        ["Expenses", -exp, "rent, salary, power, advances"],
        ["NET PROFIT", st["profit"] - exp, ""],
        ["", "", ""],
        ["CASH MOVEMENT", "", ""],
        ["Money in", cash["cash_in"], "collected from customers"],
        ["Paid to suppliers", -cash["to_suppliers"], ""],
        ["Expenses paid", -cash["expenses"], ""],
        ["Refunds", -cash["refunds"], ""],
        ["NET CASH MOVEMENT", cash["net"], ""],
        ["", "", ""],
        ["POSITION (all time)", "", ""],
        ["Cash in hand", pos["cash_in_hand"], ""],
        ["Customers owe you", pos["receivable"], ""],
        ["Stock value", pos["stock_value"], "at average cost"],
        ["You owe suppliers", -pos["payable"], ""],
        ["BUSINESS WORTH", pos["business_worth"], "cash + owed to you + stock - owed by you"],
    ]
    return {
        "title": "Overall Summary",
        "subtitle": _period_label(start, end),
        "columns": [("", "text"), ("Amount", "money"), ("Note", "text")],
        "rows": rows,
        "totals": None,
    }


REPORTS = {
    "summary": report_summary,
    "cashbook": report_cashbook,
    "sales": report_sales,
    "sale-items": report_sale_items,
    "purchases": report_purchases,
    "orders": report_orders,
    "customers": report_customers,
    "suppliers": report_suppliers,
    "items": report_items,
    "expenses": report_expenses,
    "production": report_production,
    "employees": report_employees,
}

REPORT_LABELS = {
    "summary": "Overall Summary",
    "cashbook": "Cash Book",
    "sales": "Sales",
    "sale-items": "Sales - item detail",
    "purchases": "Purchases",
    "orders": "Orders",
    "customers": "Customers & credit",
    "suppliers": "Suppliers & payables",
    "items": "Stock",
    "expenses": "Expenses",
    "production": "Production",
    "employees": "Employees",
}


def build(app, conn, name, args):
    fn = REPORTS.get(name)
    if not fn:
        return None
    return fn(app, conn, args)


# --------------------------------------------------------------------------- writers

def to_csv(rep, company=""):
    buf = io.StringIO()
    w = csv.writer(buf)
    if company:
        w.writerow([company])
    w.writerow([rep["title"]])
    w.writerow([rep["subtitle"]])
    w.writerow([])
    w.writerow([c[0] for c in rep["columns"]])
    for row in rep["rows"]:
        w.writerow(["" if v is None else v for v in row])
    if rep.get("totals"):
        w.writerow([])
        w.writerow(["" if v is None else v for v in rep["totals"]])
    return buf.getvalue()


def to_xlsx(rep, company=""):
    from openpyxl import Workbook
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

    wb = Workbook()
    ws = wb.active
    ws.title = rep["title"][:31]

    green = "2F5C3F"
    head_font = Font(bold=True, color="FFFFFF", size=10)
    head_fill = PatternFill("solid", fgColor=green)
    thin = Side(style="thin", color="D8D4CB")

    r = 1
    if company:
        ws.cell(r, 1, company).font = Font(bold=True, size=14, color=green); r += 1
    ws.cell(r, 1, rep["title"]).font = Font(bold=True, size=12); r += 1
    ws.cell(r, 1, rep["subtitle"]).font = Font(size=10, italic=True, color="6B6B63"); r += 2

    header_row = r
    for i, (label, _k) in enumerate(rep["columns"], start=1):
        c = ws.cell(r, i, label)
        c.font = head_font; c.fill = head_fill
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    r += 1

    money_fmt = '#,##0.00'
    num_fmt = '#,##0.##'
    # Build the style objects ONCE. Creating a new Border per cell made a 50,000-row
    # export take about 15 seconds; reusing them makes it a fraction of that.
    row_border = Border(bottom=thin)
    kinds = [k for _lbl, k in rep["columns"]]
    big = len(rep["rows"]) > 3000        # skip cell borders on very large exports

    for row in rep["rows"]:
        for i, val in enumerate(row, start=1):
            c = ws.cell(r, i, val)
            if not big:
                c.border = row_border
            if isinstance(val, (int, float)):
                k = kinds[i - 1] if i - 1 < len(kinds) else "text"
                if k == "money":
                    c.number_format = money_fmt
                elif k == "number":
                    c.number_format = num_fmt
        r += 1

    if rep.get("totals"):
        r += 1
        for i, (val, (_lbl, kind)) in enumerate(zip(rep["totals"], rep["columns"]), start=1):
            c = ws.cell(r, i, val)
            c.font = Font(bold=True)
            c.border = Border(top=Side(style="double", color=green))
            if kind == "money" and isinstance(val, (int, float)):
                c.number_format = money_fmt

    # Sensible column widths and a frozen header so long reports stay readable.
    for i, (label, kind) in enumerate(rep["columns"], start=1):
        longest = len(str(label))
        for row in rep["rows"][:200]:
            if i - 1 < len(row) and row[i - 1] is not None:
                longest = max(longest, len(str(row[i - 1])))
        ws.column_dimensions[ws.cell(1, i).column_letter].width = min(42, max(10, longest + 3))
    ws.freeze_panes = ws.cell(header_row + 1, 1)

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
